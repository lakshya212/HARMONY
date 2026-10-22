import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

# =============================================================================
# STUDENT & INSTITUTIONAL CONFIGURATION
# =============================================================================
STUDENT_NAME = "Lakshya Shukla"
STUDENT_FATHER_NAME = "Shri Rajiv Ranjan Shukla"
STUDENT_ROLL_NO = "240919106000"  # Format: 240919106XXX (Editable / Custom)
STUDENT_SECTION = "BCA 5th Semester, Sec: B"
STUDENT_EMAIL = "lakshyashukla.bcr24@imsnoida.com"
STUDENT_MOBILE = "+91 98XXXXXXXX"
PROJECT_TITLE = "HARMONY: ROOMMATE CONFLICT RESOLUTION AND SHARED LIVING MANAGEMENT SYSTEM"
SUBJECT_CODE = "MINOR PROJECT (BCA-508)"
UNIVERSITY_NAME = "Chaudhary Charan Singh University, Meerut"
INSTITUTION_NAME = "Institute of Management Studies (IMS), Noida"
DEPARTMENT_NAME = "School of Information Technology (SOIT)"
HOD_NAME = "Ms. Monika Dixit Bajpai"
PROJECT_GUIDE = "Mr. Nihal Kumar"
GUIDE_DESIGNATION = "Assistant Professor, SOIT, IMS Noida"
ACADEMIC_BATCH = "2024–2027"
INTERNSHIP_CERT_PATH = "plan/internship_certificate.png"
CCSU_LOGO_PATH = "plan/ccsu_logo.png"

# =============================================================================
# STYLING & XML HELPER FUNCTIONS
# =============================================================================
def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_style in borders.items():
        if border_style:
            b_el = parse_xml(f'<w:{border_name} {nsdecls("w")} w:val="{border_style.get("val", "single")}" w:sz="{border_style.get("sz", "4")}" w:space="0" w:color="{border_style.get("color", "000000")}"/>')
            tcBorders.append(b_el)
        else:
            b_el = parse_xml(f'<w:{border_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def add_page_border(section):
    sectPr = section._sectPr
    border_xml = f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page"><w:top w:val="single" w:sz="12" w:space="24" w:color="1E293B"/><w:left w:val="single" w:sz="12" w:space="24" w:color="1E293B"/><w:bottom w:val="single" w:sz="12" w:space="24" w:color="1E293B"/><w:right w:val="single" w:sz="12" w:space="24" w:color="1E293B"/></w:pgBorders>'
    sectPr.append(parse_xml(border_xml))

def add_footer_page_number(run):
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_callout_box(doc, text, title="ACADEMIC NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_borders(cell,
        top={'val': 'single', 'sz': '4', 'color': '94A3B8'},
        bottom={'val': 'single', 'sz': '4', 'color': '94A3B8'},
        left={'val': 'single', 'sz': '18', 'color': '1E3A8A'},
        right={'val': 'single', 'sz': '4', 'color': '94A3B8'}
    )
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"[{title}] ")
    run_t.bold = True
    run_t.font.color.rgb = RGBColor(30, 58, 138)
    run_t.font.size = Pt(10)
    run_m = p.add_run(text)
    run_m.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_diagram_figure(doc, image_path, caption_text, width_inches=5.8):
    if os.path.exists(image_path):
        p_pic = doc.add_paragraph()
        p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pic.paragraph_format.space_before = Pt(8)
        p_pic.paragraph_format.space_after = Pt(4)
        p_pic.add_run().add_picture(image_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

# =============================================================================
# MAIN BUILD SCRIPT
# =============================================================================
def build_complete_document():
    doc = docx.Document()
    
    # Page setup - Standard 1 inch margins with native page border
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Apply crisp academic page border to all pages
    add_page_border(section)
    
    # Configure footer page numbers (different first page for cover)
    section.different_first_page_header_footer = True
    footer = section.footer
    p_ft = footer.paragraphs[0]
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ft = p_ft.add_run()
    r_ft.font.name = 'Times New Roman'
    r_ft.font.size = Pt(10)
    r_ft.font.color.rgb = RGBColor(100, 116, 139)
    add_footer_page_number(r_ft)
    
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)
    
    # =========================================================================
    # PAGE 1: TITLE / COVER PAGE (Matches Screenshot 2026-10-06 210804.png)
    # =========================================================================
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(16)
    p1.paragraph_format.space_after = Pt(12)
    
    r_title = p1.add_run(f'“{PROJECT_TITLE}”\n')
    r_title.bold = True
    r_title.underline = True
    r_title.font.size = Pt(15.5)
    r_title.font.color.rgb = RGBColor(17, 24, 39)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(f"{SUBJECT_CODE}\nSUBMITTED TO\n")
    r_sub.bold = True
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(17, 24, 39)
    
    p_sub2 = doc.add_paragraph()
    p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub2.paragraph_format.space_after = Pt(22)
    r_sub2 = p_sub2.add_run(f"{UNIVERSITY_NAME}\n\nFor the partial fulfillment of the degree of the Bachelor’s in Computer Application\n({ACADEMIC_BATCH}).\n")
    r_sub2.bold = True
    r_sub2.font.size = Pt(12)
    
    # Centered CCSU Logo
    p_emblem = doc.add_paragraph()
    p_emblem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_emblem.paragraph_format.space_after = Pt(36)
    if os.path.exists(CCSU_LOGO_PATH):
        p_emblem.add_run().add_picture(CCSU_LOGO_PATH, width=Inches(2.2))
    
    # Bottom Guidance / Submission Table
    tbl_sub = doc.add_table(rows=1, cols=2)
    tbl_sub.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l = tbl_sub.cell(0, 0)
    cell_r = tbl_sub.cell(0, 1)
    cell_l.width = Inches(3.4)
    cell_r.width = Inches(3.1)
    
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_gl = p_l.add_run("UNDER THE GUIDANCE OF:\n")
    r_gl.bold = True
    r_gl.font.size = Pt(11)
    p_l.add_run(f"Prof. {PROJECT_GUIDE}\nIMS NOIDA")
    
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_gr = p_r.add_run("SUBMITTED BY:\n")
    r_gr.bold = True
    r_gr.font.size = Pt(11)
    p_r.add_run(f"{STUDENT_NAME.upper()}\nRoll No.: {STUDENT_ROLL_NO}\n{STUDENT_SECTION}")
    
    doc.add_page_break()
    
    # =========================================================================
    # PAGE 2: PROFORMA FOR APPROVAL (Matches Screenshot 2026-10-06 210832.png)
    # =========================================================================
    p_prof_head = doc.add_paragraph()
    p_prof_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prof_head.paragraph_format.space_before = Pt(8)
    p_prof_head.paragraph_format.space_after = Pt(12)
    r_ph1 = p_prof_head.add_run("INSTITUTE OF MANAGEMENT STUDIES, NOIDA\n")
    r_ph1.bold = True
    r_ph1.font.size = Pt(14)
    
    # Boxed Table Header
    tbl_pbox = doc.add_table(rows=1, cols=1)
    tbl_pbox.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_pb = tbl_pbox.cell(0, 0)
    cell_pb.width = Inches(6.5)
    set_cell_margins(cell_pb, top=70, bottom=70, left=120, right=120)
    set_cell_borders(cell_pb,
        top={'val': 'single', 'sz': '8', 'color': '000000'},
        bottom={'val': 'single', 'sz': '8', 'color': '000000'},
        left={'val': 'single', 'sz': '8', 'color': '000000'},
        right={'val': 'single', 'sz': '8', 'color': '000000'}
    )
    p_box_in = cell_pb.paragraphs[0]
    p_box_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pbin = p_box_in.add_run("PROFORMA FOR APPROVAL OF BCA MINOR PROJECT (BCA-508)")
    r_pbin.bold = True
    r_pbin.font.size = Pt(11.5)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Clean text fields matching Screenshot 2026-10-06 210832.png
    p_det = doc.add_paragraph()
    p_det.paragraph_format.line_spacing = 1.35
    p_det.paragraph_format.space_after = Pt(12)
    
    r = p_det.add_run("Roll no. ")
    r.bold = True
    p_det.add_run(f"{STUDENT_ROLL_NO}\n")
    
    r = p_det.add_run("Name of the Student: ")
    r.bold = True
    p_det.add_run(f"{STUDENT_NAME}\n")
    
    r = p_det.add_run("E-mail: ")
    r.bold = True
    p_det.add_run(f"{STUDENT_EMAIL}\n")
    
    r = p_det.add_run("Mob. No: ")
    r.bold = True
    p_det.add_run(f"{STUDENT_MOBILE}\n\n")
    
    r = p_det.add_run("Title of the MINOR Project: ")
    r.bold = True
    p_det.add_run(f"{PROJECT_TITLE}\n")
    
    r = p_det.add_run("Name of the Guide: ")
    r.bold = True
    p_det.add_run(f"{PROJECT_GUIDE}\n")
    
    # Horizontal line separator
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(2)
    p_line.paragraph_format.space_after = Pt(10)
    r_line = p_line.add_run("_" * 74)
    r_line.font.color.rgb = RGBColor(100, 116, 139)
    
    p_off = doc.add_paragraph()
    r_off = p_off.add_run("For Office Use Only:\n")
    r_off.bold = True
    r_off.font.size = Pt(11)
    
    # Approved / Not Approved boxes + Mentor signature block
    tbl_app = doc.add_table(rows=1, cols=3)
    tbl_app.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_app.cell(0, 0).width = Inches(2.1)
    tbl_app.cell(0, 1).width = Inches(2.1)
    tbl_app.cell(0, 2).width = Inches(2.3)
    
    c_app = tbl_app.cell(0, 0)
    set_cell_borders(c_app,
        top={'val': 'single', 'sz': '8', 'color': '000000'},
        bottom={'val': 'single', 'sz': '8', 'color': '000000'},
        left={'val': 'single', 'sz': '8', 'color': '000000'},
        right={'val': 'single', 'sz': '8', 'color': '000000'}
    )
    set_cell_margins(c_app, top=140, bottom=140, left=100, right=100)
    p_ca = c_app.paragraphs[0]
    p_ca.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ca.add_run("\n\nApproved\n")
    
    c_not = tbl_app.cell(0, 1)
    set_cell_borders(c_not,
        top={'val': 'single', 'sz': '8', 'color': '000000'},
        bottom={'val': 'single', 'sz': '8', 'color': '000000'},
        left={'val': 'single', 'sz': '8', 'color': '000000'},
        right={'val': 'single', 'sz': '8', 'color': '000000'}
    )
    set_cell_margins(c_not, top=140, bottom=140, left=100, right=100)
    p_cn = c_not.paragraphs[0]
    p_cn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cn.add_run("\n\nNot Approved\n")
    
    c_sig = tbl_app.cell(0, 2)
    p_cs = c_sig.paragraphs[0]
    p_cs.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_cs.add_run("\n\nSignature of the Mentor\n\nDate: _____ / _____ / 2026").bold = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    p_sug_head = doc.add_paragraph()
    r_sh = p_sug_head.add_run("Suggestions (if any): -\n")
    r_sh.bold = True
    r_sh.font.size = Pt(11)
    
    for num in range(1, 6):
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_after = Pt(8)
        p_s.add_run(f"{num}. _____________________________________________________________________")
        
    doc.add_page_break()
    
    # =========================================================================
    # PAGE 3: ACKNOWLEDGEMENT (Matches Screenshot 2026-10-06 210814.png)
    # =========================================================================
    p_ack_h = doc.add_paragraph()
    p_ack_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack_h.paragraph_format.space_before = Pt(14)
    p_ack_h.paragraph_format.space_after = Pt(20)
    r_ah = p_ack_h.add_run("ACKNOWLEDGEMENT\n")
    r_ah.bold = True
    r_ah.underline = True
    r_ah.font.size = Pt(15)
    
    p_ack1 = doc.add_paragraph()
    p_ack1.paragraph_format.line_spacing = 1.25
    p_ack1.paragraph_format.space_after = Pt(14)
    p_ack1.add_run(
        f"I would like to express my sincere gratitude to my project mentor, Prof. {PROJECT_GUIDE}, "
        f"for his valuable guidance, continuous support, and encouragement throughout the development of my project titled "
        f"\"{PROJECT_TITLE}.\" His expertise and suggestions greatly contributed to the successful completion of this work."
    )
    
    p_ack2 = doc.add_paragraph()
    p_ack2.paragraph_format.line_spacing = 1.25
    p_ack2.paragraph_format.space_after = Pt(14)
    p_ack2.add_run(
        f"I would also like to express my sincere thanks to our respected Head of Department (HOD SOIT), "
        f"{HOD_NAME}, for providing an academically stimulating environment, constant encouragement, "
        "and administrative facilities essential for the development of our Minor Project."
    )
    
    p_ack3 = doc.add_paragraph()
    p_ack3.paragraph_format.line_spacing = 1.25
    p_ack3.paragraph_format.space_after = Pt(14)
    p_ack3.add_run(
        "I am also thankful to the faculty members and the Department of School of Information Technology (SOIT), "
        "IMS Noida, for providing the necessary facilities, computing resources, and academic support required for this project."
    )
    
    p_ack4 = doc.add_paragraph()
    p_ack4.paragraph_format.line_spacing = 1.25
    p_ack4.paragraph_format.space_after = Pt(14)
    p_ack4.add_run(
        "I extend my heartfelt thanks to my family and friends for their constant motivation, patience, and encouragement during the course of this project."
    )
    
    p_ack5 = doc.add_paragraph()
    p_ack5.paragraph_format.line_spacing = 1.25
    p_ack5.paragraph_format.space_after = Pt(40)
    p_ack5.add_run(
        "Finally, I would like to thank everyone who directly or indirectly contributed to the successful completion of this project."
    )
    
    # Signature block matching Screenshot 2026-10-06 210814.png
    p_ack_sig = doc.add_paragraph()
    p_ack_sig.paragraph_format.space_after = Pt(8)
    p_ack_sig.add_run(f"{STUDENT_NAME}-                  {STUDENT_ROLL_NO}\n").bold = True
    p_ack_sig.add_run("Signature-                       ............................")
    
    doc.add_page_break()
    
    # =========================================================================
    # PAGE 4: CERTIFICATE OF ORIGINALITY (Matches Screenshot 2026-10-06 210821.png)
    # =========================================================================
    p_cert_h = doc.add_paragraph()
    p_cert_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_h.paragraph_format.space_before = Pt(14)
    p_cert_h.paragraph_format.space_after = Pt(20)
    r_ch = p_cert_h.add_run("CERTIFICATE OF ORIGINALITY\n")
    r_ch.bold = True
    r_ch.font.size = Pt(15)
    
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.line_spacing = 1.25
    p_c1.paragraph_format.space_after = Pt(16)
    p_c1.add_run(
        f"This is to certify that the project report entitled \"{PROJECT_TITLE}\" "
        f"submitted by {STUDENT_NAME} in partial fulfillment of the requirements for the award of the degree of "
        f"Bachelor of Computer Applications (BCA) to {UNIVERSITY_NAME}, through {INSTITUTION_NAME}, "
        f"is an original work carried out by the student under the guidance and supervision of Prof. {PROJECT_GUIDE}."
    )
    
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.line_spacing = 1.25
    p_c2.paragraph_format.space_after = Pt(16)
    p_c2.add_run(
        "The work embodied in this project report has not been submitted, either in part or in full, "
        "to any other University or Institution for the award of any degree, diploma, or other academic qualification."
    )
    
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.line_spacing = 1.25
    p_c3.paragraph_format.space_after = Pt(40)
    p_c3.add_run(
        f"The project is a bona fide record of the work carried out by the student during the academic session {ACADEMIC_BATCH}."
    )
    
    tbl_csig = doc.add_table(rows=1, cols=2)
    tbl_csig.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_g = tbl_csig.cell(0, 0)
    c_s = tbl_csig.cell(0, 1)
    c_g.width = Inches(3.4)
    c_s.width = Inches(3.1)
    
    p_cg = c_g.paragraphs[0]
    p_cg.add_run("Project Guide\n").bold = True
    p_cg.add_run(f"Prof. {PROJECT_GUIDE}\nDepartment of SOIT\nIMS NOIDA\n\nDate: ........................\nPlace: Noida")
    
    p_cs = c_s.paragraphs[0]
    p_cs.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_cs.add_run("Student\n").bold = True
    p_cs.add_run(f"{STUDENT_NAME}\nBCA 5th Semester\n{STUDENT_ROLL_NO}\n\nSignature: ........................")
    
    doc.add_page_break()
    
    # =========================================================================
    # PAGE 5: MANDATORY INTERNSHIP COMPLETION CERTIFICATE (Matches Screenshot 2026-10-06 210826.png)
    # =========================================================================
    p_int_pic = doc.add_paragraph()
    p_int_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_int_pic.paragraph_format.space_before = Pt(8)
    p_int_pic.paragraph_format.space_after = Pt(10)
    if os.path.exists(INTERNSHIP_CERT_PATH):
        p_int_pic.add_run().add_picture(INTERNSHIP_CERT_PATH, width=Inches(6.3))
    
    p_int_foot = doc.add_paragraph()
    p_int_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_int_foot.paragraph_format.space_after = Pt(4)
    r_if1 = p_int_foot.add_run("Official Internship Certificate | Certificate No: CPL-INT-2026-012 | Cannock Private Limited (CIN: U72900DL2021PTC389998)\n")
    r_if1.font.size = Pt(9.5)
    r_if1.font.italic = True
    r_if1.font.color.rgb = RGBColor(71, 85, 105)
    
    r_if2 = p_int_foot.add_run(f"Student: {STUDENT_NAME} (S/o {STUDENT_FATHER_NAME}) | Domain: AI & Automation | Duration: 13 June 2026 – 13 August 2026 | Verified: office@cannock.in")
    r_if2.font.size = Pt(9)
    r_if2.font.bold = True
    r_if2.font.color.rgb = RGBColor(30, 58, 138)
    
    doc.add_page_break()
    
    # =========================================================================
    # PAGE 6: TABLE OF CONTENTS / INDEX (Matches Screenshot 2026-10-06 210838.png & 210844.png)
    # =========================================================================
    p_idx_h = doc.add_paragraph()
    p_idx_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_idx_h.paragraph_format.space_before = Pt(8)
    p_idx_h.paragraph_format.space_after = Pt(14)
    r_idx = p_idx_h.add_run("TABLE OF CONTENTS / INDEX\n")
    r_idx.bold = True
    r_idx.font.size = Pt(15)
    
    index_rows = [
        ("•", "Proforma for Approval of BCA Minor Project", "2"),
        ("•", "Acknowledgement", "3"),
        ("•", "Certificate of Originality", "4"),
        ("•", "Internship Completion Certificate (Cannock Pvt Ltd)", "5"),
        ("•", "List of Figures", "7"),
        ("1", "Introduction / Overview of the Project", "8"),
        ("   1.1", "Problem Background & Problem Statement", "8"),
        ("   1.2", "Real-World Relevance & Target Demographics", "9"),
        ("   1.3", "Proposed Solution: The HARMONY System", "9"),
        ("   1.4", "Primary System Objectives", "10"),
        ("2", "Performance of the Project", "11"),
        ("   2.1", "Feasibility Analysis (Technical, Operational, Economic, Scheduling)", "11"),
        ("   2.2", "System Effectiveness & Operational Metrics", "12"),
        ("   2.3", "Scalability, Latency & Capacity Benchmarks", "13"),
        ("3", "Project Architecture", "14"),
        ("   3.1", "3-Tier Decoupled Web Architecture & Communication Pipeline", "14"),
        ("   3.2", "Presentation Tier (React 19 + Tailwind CSS + Lucide Icons)", "15"),
        ("   3.3", "Application / Logic Tier (Node.js + Express REST API)", "16"),
        ("   3.4", "Data Tier (PostgreSQL + Supabase Cloud Engine)", "16"),
        ("   3.5", "Security & Data Flow Pipeline", "17"),
        ("4", "Project Flow", "18"),
        ("   4.1", "End-to-End System Flowchart & Lifecycle Analysis", "18"),
        ("   4.2", "Step-by-Step Functional Workflow & State Transitions", "19"),
        ("5", "Data Flow Diagrams (DFDs) [Each on a Separate Page]", "20"),
        ("   5.1", "Level 0 DFD (Context Diagram) [Separate Page]", "21"),
        ("   5.2", "Level 1 DFD (Sub-System Functional Decomposition) [Separate Page]", "22"),
        ("   5.3", "Level 2 DFD (Transaction & Calculation Processing) [Separate Page]", "23"),
        ("6", "Use Case Diagram", "24"),
        ("   6.1", "System Actors & Role Identification", "24"),
        ("   6.2", "Unified UML Use Case Diagram & Specifications", "25"),
        ("7", "Entity Relationship Diagram (ERD)", "26"),
        ("   7.1", "Relational Entity Specifications & Domain Model", "26"),
        ("   7.2", "ER Diagram Schema & Relational Integrity Rules", "27"),
        ("8", "Hardware and Software Requirements", "29"),
        ("   8.1", "Development & Production Hardware Specifications", "29"),
        ("   8.2", "Software & Technology Stack Specifications", "30"),
        ("   8.3", "Cloud Infrastructure & ₹0 Zero-Cost Deployment Stack", "31"),
        ("9", "Future Scope", "32"),
        ("   9.1", "Machine Learning Fair Share Chore Optimizer", "32"),
        ("   9.2", "UPI Deep-Linking & Auto-Payment Reconciliation", "32"),
        ("   9.3", "Native Mobile Application & Push Notifications", "33"),
        ("10", "Limitations of the Project", "34"),
        ("   10.1", "Technical & Operational Constraints", "34"),
        ("   10.2", "Mitigation Strategies & System Safeguards", "35"),
        ("11", "Front-End Pages (UI Mockups & Screen Layouts)", "36"),
        ("   11.1", "Modern Tech Startup Landing Page (Zero Academic Clutter)", "36"),
        ("   11.2", "Interactive Household Dashboard & Net Balance Matrix", "37"),
        ("   11.3", "Split Expense Modal & Real-Time Calculation Engine", "38"),
        ("   11.4", "User Profile & Household Administrative Governance Console", "38"),
        ("•", "References & Bibliography", "39")
    ]
    
    tbl_idx = doc.add_table(rows=len(index_rows) + 1, cols=3)
    tbl_idx.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_idx.cell(0, 0).width = Inches(0.8)
    tbl_idx.cell(0, 1).width = Inches(4.7)
    tbl_idx.cell(0, 2).width = Inches(1.0)
    
    idx_headers = ["S.No.", "Topic / Chapter Title", "Page No."]
    for c_i, h_text in enumerate(idx_headers):
        c = tbl_idx.cell(0, c_i)
        set_cell_background(c, "E2E8F0")
        set_cell_margins(c, 80, 80, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(11)
        if c_i == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
    for r_i, (sno, title, pageno) in enumerate(index_rows, start=1):
        c0 = tbl_idx.cell(r_i, 0)
        c1 = tbl_idx.cell(r_i, 1)
        c2 = tbl_idx.cell(r_i, 2)
        set_cell_margins(c0, 50, 50, 80, 80)
        set_cell_margins(c1, 50, 50, 80, 80)
        set_cell_margins(c2, 50, 50, 80, 80)
        c0.paragraphs[0].add_run(sno).font.size = Pt(10)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(title)
        r1.font.size = Pt(10)
        if sno.strip() in ["1","2","3","4","5","6","7","8","9","10","11"]:
            r1.bold = True
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.add_run(pageno).font.size = Pt(10)
        
    doc.add_page_break()
    
    # =========================================================================
    # PAGE 7: LIST OF FIGURES
    # =========================================================================
    p_fig_h = doc.add_paragraph()
    p_fig_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fig_h.paragraph_format.space_before = Pt(8)
    p_fig_h.paragraph_format.space_after = Pt(14)
    r_figh = p_fig_h.add_run("LIST OF FIGURES\n")
    r_figh.bold = True
    r_figh.font.size = Pt(15)
    
    figures_list = [
        ("Figure 3.1", "HARMONY 3-Tier Web Architecture Diagram", "14"),
        ("Figure 4.1", "End-to-End System Interaction & State Flowchart", "18"),
        ("Figure 5.1", "Level 0 DFD - Context Diagram (Boundary Model)", "21"),
        ("Figure 5.2", "Level 1 DFD - Sub-System Functional Decomposition", "22"),
        ("Figure 5.3", "Level 2 DFD - Financial Ledger & Debt Settlement Processing", "23"),
        ("Figure 6.1", "Unified UML Use Case Diagram", "25"),
        ("Figure 7.1", "Relational Entity-Relationship Diagram (Crow's Foot ERD)", "27"),
        ("Figure 11.1", "Modern Startup Landing Page Interface (Zero Academic Clutter)", "36"),
        ("Figure 11.2", "Interactive Household Dashboard & Net Balance Matrix", "37"),
        ("Figure 11.3", "Split Expense Modal Interface & Live Multi-User Calculation", "38"),
        ("Figure 11.4", "Administrative Household Supervision & Governance Panel", "38")
    ]
    
    tbl_fig = doc.add_table(rows=len(figures_list) + 1, cols=3)
    tbl_fig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_fig.cell(0, 0).width = Inches(1.2)
    tbl_fig.cell(0, 1).width = Inches(4.3)
    tbl_fig.cell(0, 2).width = Inches(1.0)
    
    for c_i, h_text in enumerate(["Figure No.", "Figure Title & Description", "Page No."]):
        c = tbl_fig.cell(0, c_i)
        set_cell_background(c, "E2E8F0")
        set_cell_margins(c, 80, 80, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(11)
        if c_i == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
    for r_i, (fno, ftitle, fpno) in enumerate(figures_list, start=1):
        c0 = tbl_fig.cell(r_i, 0)
        c1 = tbl_fig.cell(r_i, 1)
        c2 = tbl_fig.cell(r_i, 2)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        set_cell_margins(c2, 60, 60, 80, 80)
        c0.paragraphs[0].add_run(fno).bold = True
        c1.paragraphs[0].add_run(ftitle)
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.add_run(fpno)
        
    doc.add_page_break()
    
    return doc

print("Front matter generator ready.")

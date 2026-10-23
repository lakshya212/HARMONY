import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(15, 23, 42)        # #0F172A Deep Slate Navy
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B Slate Container
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155 Border
    TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    ACCENT_EMERALD = RGBColor(16, 185, 129) # #10B981 Emerald
    ACCENT_VIOLET = RGBColor(139, 92, 246)  # #8B5CF6 Violet
    ACCENT_AMBER = RGBColor(245, 158, 11)   # #F59E0B Amber
    ACCENT_CYAN = RGBColor(6, 182, 212)     # #06B6D4 Cyan

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="BCA-508 MINOR PROJECT FINAL PRESENTATION"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_EMERALD

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.fill.background()

    def add_footer(slide, current_slide, total_slides=12):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.35))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"HARMONY | Candidate: Lakshya Shukla (240919106121) | Guide: Mr. Nihal Kumar | IMS Noida (Panel 2 - Room DG4) | Slide {current_slide} of {total_slides}"
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title=None, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = ACCENT_CYAN
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE (WITH CLICKABLE LINKS)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.55), Inches(3.2), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_EMERALD
    bar.line.fill.background()

    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(2.25))
    tf = title_box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "HARMONY"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE

    p1 = tf.add_paragraph()
    p1.text = "Shared Living Management & Conflict Resolution Platform"
    p1.font.size = Pt(21)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_EMERALD

    p2 = tf.add_paragraph()
    p2.text = "Minor Project Final Presentation (BCA-508) | Full-Stack System for Expense Splitting, Effort Gamification & Household Transparency"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED

    # Academic info card
    add_card(s1, Inches(0.8), Inches(3.2), Inches(5.6), Inches(3.45), title="ACADEMIC CREDENTIALS")
    tb_meta = s1.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(5.2), Inches(2.8))
    tf_m = tb_meta.text_frame
    tf_m.word_wrap = True

    def add_meta_field(tf_obj, label, value, val_color=TEXT_WHITE):
        p = tf_obj.add_paragraph() if tf_obj.paragraphs[0].text else tf_obj.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{label}: "
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = ACCENT_CYAN

        r2 = p.add_run()
        r2.text = f"{value}\n"
        r2.font.size = Pt(11)
        r2.font.color.rgb = val_color

    add_meta_field(tf_m, "Course & Code", "BCA 5th Semester | Minor Project (BCA-508)")
    add_meta_field(tf_m, "Candidate Name", "Lakshya Shukla")
    add_meta_field(tf_m, "University Roll No", "240919106121 (Section B)")
    add_meta_field(tf_m, "College Email", "lakshyashukla.bcr24257@imsnoida.com")
    add_meta_field(tf_m, "Project Guide", "Mr. Nihal Kumar, Assistant Professor")
    add_meta_field(tf_m, "Review Panel", "Panel 2 (Room DG4) | Final Presentation Day")
    add_meta_field(tf_m, "Institution", "IMS Noida | Affiliated to CCS University, Meerut")

    # Mandatory Live Links Card
    add_card(s1, Inches(6.7), Inches(3.2), Inches(5.8), Inches(3.45), title="MANDATORY FINAL PRESENTATION LIVE LINKS", border_color=ACCENT_EMERALD)
    tb_links = s1.shapes.add_textbox(Inches(6.9), Inches(3.7), Inches(5.4), Inches(2.8))
    tf_l = tb_links.text_frame
    tf_l.word_wrap = True

    pl0 = tf_l.paragraphs[0]
    pl0.text = "Clickable Production Access (As Mandated by Department):"
    pl0.font.size = Pt(11)
    pl0.font.italic = True
    pl0.font.color.rgb = TEXT_MUTED

    pl1 = tf_l.add_paragraph()
    r = pl1.add_run()
    r.text = "\n[LIVE PRODUCTION DEPLOYMENT] "
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ACCENT_EMERALD

    r_url = pl1.add_run()
    r_url.text = "Click to Open Live HARMONY Web App\n"
    r_url.font.size = Pt(13)
    r_url.font.bold = True
    r_url.font.underline = True
    r_url.font.color.rgb = ACCENT_CYAN
    r_url.hyperlink.address = "https://harmony-household.vercel.app"

    pl2 = tf_l.add_paragraph()
    r2 = pl2.add_run()
    r2.text = "[PUBLIC GITHUB REPOSITORY & CODE] "
    r2.font.size = Pt(12)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT_VIOLET

    r2_url = pl2.add_run()
    r2_url.text = "github.com/lakshya212/HARMONY\n"
    r2_url.font.size = Pt(13)
    r2_url.font.bold = True
    r2_url.font.underline = True
    r2_url.font.color.rgb = ACCENT_CYAN
    r2_url.hyperlink.address = "https://github.com/lakshya212/HARMONY"

    pl3 = tf_l.add_paragraph()
    r3 = pl3.add_run()
    r3.text = "Architecture: React 18 + Node.js/Express + PostgreSQL\nOffline-Resilient Failover Layer Active for Zero-Downtime Defense."
    r3.font.size = Pt(10)
    r3.font.color.rgb = ACCENT_AMBER

    add_footer(s1, 1)

    # ==========================================
    # SLIDE 2: PROJECT OVERVIEW & MOTIVATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, "Project Overview & Problem Statement")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.9), title="THE PROBLEM CONTEXT")
    tb2_1 = s2.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(3.2), Inches(4.1))
    tf2_1 = tb2_1.text_frame
    tf2_1.word_wrap = True
    p = tf2_1.paragraphs[0]
    p.text = "Shared Household Friction:\n\n• Asymmetric Financial Burden: Flatmates and multi-earner households frequently lose track of who paid for groceries, utilities, and internet bills, causing awkward disputes.\n\n• The 'Chore Trap': Traditional chore tracking ignores effort disparity (e.g., taking out trash vs deep cleaning bathroom treated equally).\n\n• Unreliable Memory: Informal verbal agreements lead to resentment and roommate conflicts."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.9), title="THE HARMONY SOLUTION")
    tb2_2 = s2.shapes.add_textbox(Inches(5.0), Inches(2.4), Inches(3.2), Inches(4.1))
    tf2_2 = tb2_2.text_frame
    tf2_2.word_wrap = True
    p = tf2_2.paragraphs[0]
    p.text = "Mathematical Transparency:\n\n• Automated Split Engine: Instantly splits shared expenses equally or by shares, updating a real-time Net Balance Ledger.\n\n• Gamified Effort Points: Chores carry calibrated effort points (10-30 pts). System highlights top contributors dynamically.\n\n• Unified Dashboard: Gives roommates an instant high-level view of who owes what and who completed pending duties.\n\n• Zero Setup Overhead: Instant household code generation for roommates."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.9), title="ACADEMIC OBJECTIVES")
    tb2_3 = s2.shapes.add_textbox(Inches(9.0), Inches(2.4), Inches(3.3), Inches(4.1))
    tf2_3 = tb2_3.text_frame
    tf2_3.word_wrap = True
    p = tf2_3.paragraphs[0]
    p.text = "BCA-508 Minor Project Scope:\n\n• Full-Stack Architecture: Modern React Single Page Application (SPA) consuming decoupled Node.js RESTful APIs.\n\n• Relational Integrity: Enforcing 3NF relational schemas with Foreign Keys and transactional integrity in PostgreSQL.\n\n• Production Cloud Deployment: Automated CD workflow with Vercel Edge hosting and cloud backend services.\n\n• Defensive Reliability: Dual-mode data resilience preventing presentation failure during network outages."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s2, 2)

    # ==========================================
    # SLIDE 3: SYSTEM ARCHITECTURE & TECH STACK
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, "Three-Tier System Architecture & Tech Stack")

    arch_img = "diagrams/diagram_architecture.png"
    if os.path.exists(arch_img):
        s3.shapes.add_picture(arch_img, Inches(0.8), Inches(1.8), Inches(6.2), Inches(4.9))

    add_card(s3, Inches(7.3), Inches(1.8), Inches(5.2), Inches(4.9), title="TECHNICAL IMPLEMENTATION")
    tb3 = s3.shapes.add_textbox(Inches(7.5), Inches(2.4), Inches(4.8), Inches(4.1))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Presentation Tier (Frontend):\n• React 18 with Vite build tooling for sub-second HMR.\n• Tailwind CSS for modern responsive dark-mode styling.\n• Lucide React icon library and client-side SPA routing.\n\nApplication Tier (Backend API):\n• Node.js & Express.js RESTful API architecture.\n• Middleware: CORS protection, JSON body parsing, and request logging.\n• Strict route isolation: /api/auth, /api/chores, /api/expenses, /api/household.\n\nData Tier (Persistence):\n• PostgreSQL relational database (hosted on Supabase/Neon).\n• Tables: Users, Households, Household_Members, Chores, Expenses, Expense_Splits.\n• Failover Resilience: Seamless fallback to client-side synchronized memory store if remote host is cold-starting."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s3, 3)

    # ==========================================
    # SLIDE 4: DATA FLOW DIAGRAMS (DFD 0 & 1)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, "Data Flow Diagrams (Context Level 0 & Level 1)")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), title="DFD LEVEL 0: CONTEXT MODEL")
    dfd0_img = "diagrams/diagram_dfd0.png"
    if os.path.exists(dfd0_img):
        s4.shapes.add_picture(dfd0_img, Inches(1.0), Inches(2.4), Inches(5.2), Inches(3.2))
    tb_dfd0 = s4.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(5.2), Inches(0.8))
    tf_dfd0 = tb_dfd0.text_frame
    tf_dfd0.word_wrap = True
    p = tf_dfd0.paragraphs[0]
    p.text = "Identifies boundary: External entities (Roommate, Household Admin) interact with HARMONY 0.0 System Boundary via secure HTTPS requests."
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED

    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9), title="DFD LEVEL 1: FUNCTIONAL DECOMPOSITION")
    dfd1_img = "diagrams/diagram_dfd1.png"
    if os.path.exists(dfd1_img):
        s4.shapes.add_picture(dfd1_img, Inches(7.0), Inches(2.4), Inches(5.3), Inches(3.2))
    tb_dfd1 = s4.shapes.add_textbox(Inches(7.0), Inches(5.7), Inches(5.3), Inches(0.8))
    tf_dfd1 = tb_dfd1.text_frame
    tf_dfd1.word_wrap = True
    p = tf_dfd1.paragraphs[0]
    p.text = "Processes: 1.0 Auth & Profile | 2.0 Household Management | 3.0 Chore Task Engine | 4.0 Shared Expense Splitting | 5.0 Net Settlement & Balance."
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED

    add_footer(s4, 4)

    # ==========================================
    # SLIDE 5: DATABASE SCHEMA & ER DIAGRAM
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, "Relational Database Design & Entity Relationships")

    erd_img = "diagrams/diagram_erd.png"
    if os.path.exists(erd_img):
        s5.shapes.add_picture(erd_img, Inches(0.8), Inches(1.8), Inches(6.2), Inches(4.9))

    add_card(s5, Inches(7.3), Inches(1.8), Inches(5.2), Inches(4.9), title="SCHEMA STRUCTURE & INTEGRITY")
    tb5 = s5.shapes.add_textbox(Inches(7.5), Inches(2.4), Inches(4.8), Inches(4.1))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    p = tf5.paragraphs[0]
    p.text = "Relational Entities (PostgreSQL 3NF):\n\n1. users:\n   UUID primary key, email (unique), full_name, role.\n\n2. households:\n   id, household_name, invite_code (unique index), created_by.\n\n3. household_members (Junction):\n   household_id (FK), user_id (FK), joined_at, role (Admin/Member).\n\n4. chores:\n   id, household_id (FK), title, points (10-30), assigned_to (FK), status (pending/completed).\n\n5. expenses:\n   id, household_id (FK), title, amount, category, payer_id (FK), date.\n\n6. expense_splits:\n   expense_id (FK), user_id (FK), owed_amount, is_settled.\n\nReferential Integrity: ON DELETE CASCADE applied to household joins to prevent orphan ledger records."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s5, 5)

    # ==========================================
    # SLIDE 6: CORE FEATURE 1 - SMART EXPENSE SPLITTING
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, "Core Module 1: Automated Expense Splitting & Ledger")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), title="MATHEMATICAL SPLIT FORMULATION")
    tb6_1 = s6.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.1))
    tf6_1 = tb6_1.text_frame
    tf6_1.word_wrap = True
    p = tf6_1.paragraphs[0]
    p.text = "Mathematical Principle of Ledger Equilibrium:\n\nFor any household member i:\n   Net Balance (B_i) = Total Amount Paid (P_i) - Expected Share (S_i)\n\nEquilibrium Invariant:\n   Sum of all B_i across household = 0\n\n• Positive Net Balance (+₹X):\n  The member has paid more than their equal share. The group owes this member ₹X. Displayed in Emerald Green.\n\n• Negative Net Balance (-₹Y):\n  The member has undercontributed or paid less than their share. This member owes ₹Y to the group. Displayed in Amber/Rose.\n\n• Zero Net Balance (₹0):\n  The member's dues are perfectly settled."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9), title="EXPENSE ENGINE CAPABILITIES")
    tb6_2 = s6.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.1))
    tf6_2 = tb6_2.text_frame
    tf6_2.word_wrap = True
    p = tf6_2.paragraphs[0]
    p.text = "Operational Workflow:\n\n1. Bill Recording:\n   Roommate records expense (e.g. ₹2,400 for Fiber Internet, paid by Lakshya Shukla).\n\n2. Split Computation:\n   System divides amount equally across all 4 household occupants: ₹600 per occupant.\n\n3. Ledger Update:\n   • Lakshya Paid: ₹2,400 | Share: ₹600 | Net Balance: +₹1,800\n   • Other 3 Roommates: Paid: ₹0 | Share: ₹600 | Net Balance: -₹600 each\n\n4. Transparent History:\n   Categorized under Groceries, Utilities, Internet, Supplies with exact timestamps and payer tags."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s6, 6)

    # ==========================================
    # SLIDE 7: CORE FEATURE 2 - EFFORT POINT GAMIFICATION
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    add_header(s7, "Core Module 2: Effort-Weighted Chore Allocation")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), title="THE EFFORT POINT SYSTEM")
    tb7_1 = s7.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.1))
    tf7_1 = tb7_1.text_frame
    tf7_1.word_wrap = True
    p = tf7_1.paragraphs[0]
    p.text = "Eliminating Subjectivity in Household Duties:\n\nIn standard flatmates setups, simple binary chore counts treat 5 minutes of wiping a table the same as 1 hour of deep cleaning bathroom tiles.\n\nHARMONY Effort Matrix:\n• Light Task (10 Points): Take out daily trash, sweep hallway.\n• Medium Task (20 Points): Wash shared kitchen dishes, grocery run.\n• Heavy Task (30 Points): Deep clean bathroom, mop living area.\n\nDynamic Leaderboard:\nUpon marking a chore 'Completed', the points are awarded to the responsible member. A visual effort distribution bar displays the true contribution percentage."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9), title="DUAL-METRIC BALANCE")
    tb7_2 = s7.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.1))
    tf7_2 = tb7_2.text_frame
    tf7_2.word_wrap = True
    p = tf7_2.paragraphs[0]
    p.text = "Holistic Roommate Equity:\n\nHARMONY reconciles both sides of shared living:\n\n1. Financial Contribution (Money):\n   Tracks rupee expenses to avoid monetary exploitation.\n\n2. Physical Contribution (Labor):\n   Tracks effort points to avoid chore freeloading.\n\nBalanced Household Example:\nIf Flatmate A earns less but takes on 45% of high-point chores, while Flatmate B works late hours and covers utilities, both flatmates have complete, transparent visibility without emotional tension."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s7, 7)

    # ==========================================
    # SLIDE 8: UI WALKTHROUGH & USER EXPERIENCE
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    add_header(s8, "Modern UI Design & Interactive Interface")

    landing_img = "diagrams/ui_landing.png"
    if os.path.exists(landing_img):
        s8.shapes.add_picture(landing_img, Inches(0.8), Inches(1.8), Inches(5.6), Inches(3.2))

    dash_img = "diagrams/ui_dashboard.png"
    if os.path.exists(dash_img):
        s8.shapes.add_picture(dash_img, Inches(6.8), Inches(1.8), Inches(5.7), Inches(3.2))

    add_card(s8, Inches(0.8), Inches(5.2), Inches(5.6), Inches(1.6), title="LANDING PAGE & ONBOARDING")
    tb8_1 = s8.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(5.2), Inches(1.1))
    tf8_1 = tb8_1.text_frame
    tf8_1.word_wrap = True
    p = tf8_1.paragraphs[0]
    p.text = "Clean, professional dark-mode design highlighting core value propositions, feature highlights, and frictionless authentication."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_WHITE

    add_card(s8, Inches(6.8), Inches(5.2), Inches(5.7), Inches(1.6), title="ROOMMATE LIVE DASHBOARD")
    tb8_2 = s8.shapes.add_textbox(Inches(7.0), Inches(5.6), Inches(5.3), Inches(1.1))
    tf8_2 = tb8_2.text_frame
    tf8_2.word_wrap = True
    p = tf8_2.paragraphs[0]
    p.text = "Real-time Net Balance card, Effort Leaderboard, Pending vs Completed chores, and recent shared bills with one-click actions."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s8, 8)

    # ==========================================
    # SLIDE 9: ADMIN CONTROL & SECURITY (RBAC)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)
    add_header(s9, "Role-Based Access Control (RBAC) & Governance")

    admin_img = "diagrams/ui_admin.png"
    if os.path.exists(admin_img):
        s9.shapes.add_picture(admin_img, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.9))

    add_card(s9, Inches(7.1), Inches(1.8), Inches(5.4), Inches(4.9), title="GOVERNANCE & PERMISSIONS")
    tb9 = s9.shapes.add_textbox(Inches(7.3), Inches(2.4), Inches(5.0), Inches(4.1))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    p = tf9.paragraphs[0]
    p.text = "Role Hierarchy:\n\n• Household Admin (Primary Householder):\n  Can create households, issue invite codes, assign chores, remove flatmates, and settle dispute adjustments.\n\n• Household Member:\n  Can view dashboard, log new bills, claim/complete chores, and inspect net balances.\n\n• System SuperAdmin:\n  Full audit visibility across all households, system health, and database purge tools.\n\nSecurity Protections:\n• Input sanitization against XSS and SQL injection.\n• JWT / Bearer Token authorization verification.\n• Strict household tenancy isolation (members cannot view expenses of other households)."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s9, 9)

    # ==========================================
    # SLIDE 10: TESTING, VERIFICATION & RELIABILITY
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_background(s10)
    add_header(s10, "Quality Assurance, Testing & Presentation Defense")

    add_card(s10, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.9), title="TEST SUITE EXECUTION")
    tb10_1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(3.2), Inches(4.1))
    tf10_1 = tb10_1.text_frame
    tf10_1.word_wrap = True
    p = tf10_1.paragraphs[0]
    p.text = "Verified Test Cases:\n\n• Auth Verification:\n  Valid login, wrong credentials, session token persistence.\n\n• Ledger Verification:\n  Zero-sum test: Sum of all net balances strictly equals 0.\n\n• Task Life Cycle:\n  Pending -> In-Progress -> Completed -> Point Increment.\n\n• Edge-Case Handling:\n  Zero-amount expense rejected; duplicate member prevention."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_WHITE

    add_card(s10, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.9), title="OFFLINE RESILIENCE")
    tb10_2 = s10.shapes.add_textbox(Inches(5.0), Inches(2.4), Inches(3.2), Inches(4.1))
    tf10_2 = tb10_2.text_frame
    tf10_2.word_wrap = True
    p = tf10_2.paragraphs[0]
    p.text = "Zero-Downtime Defense:\n\nFree cloud services (Render, Supabase) may experience a 30-50 second cold-start spin-up.\n\nHARMONY Dual-Store Innovation:\nThe frontend React client checks API reachability. If the remote API is warming up, the UI automatically falls back to an authenticated synchronized memory store.\n\nResult: 0% chance of presentation crash or blank page during panel inspection!"
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_WHITE

    add_card(s10, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.9), title="VIVA DEFENSE READINESS")
    tb10_3 = s10.shapes.add_textbox(Inches(9.0), Inches(2.4), Inches(3.3), Inches(4.1))
    tf10_3 = tb10_3.text_frame
    tf10_3.word_wrap = True
    p = tf10_3.paragraphs[0]
    p.text = "Key Concepts for Panel 2:\n\n• 3-Tier Architecture:\n  Clear decoupling between React client, Express API, and SQL DB.\n\n• RESTful Methods:\n  GET (retrieve), POST (create), PUT (update state), DELETE (remove).\n\n• Normalization:\n  Eliminating insertion, deletion, and update anomalies via 3NF tables.\n\n• State Management:\n  React Context API + Hooks for predictable frontend data flow."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_WHITE

    add_footer(s10, 10)

    # ==========================================
    # SLIDE 11: DEPLOYMENT ARCHITECTURE & LIVE ACCESS
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_slide_background(s11)
    add_header(s11, "Cloud Deployment Topology & Mandatory Live Links")

    add_card(s11, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.9), title="PRODUCTION CLOUD DEPLOYMENT TOPOLOGY", border_color=ACCENT_EMERALD)
    tb11 = s11.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(11.1), Inches(4.1))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "Deployment Architecture Breakdown (Mandated by Department for Final Presentation):"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD

    p_c1 = tf11.add_paragraph()
    r = p_c1.add_run()
    r.text = "\n1. Frontend SPA Deployment: Vercel Edge Network\n"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ACCENT_CYAN

    r_desc1 = p_c1.add_run()
    r_desc1.text = "   • Built with Vite, bundled to /dist, distributed globally across CDN edge points.\n   • Direct Clickable Live URL: "
    r_desc1.font.size = Pt(11)
    r_desc1.font.color.rgb = TEXT_WHITE

    r_link1 = p_c1.add_run()
    r_link1.text = "https://harmony-household.vercel.app\n"
    r_link1.font.size = Pt(12)
    r_link1.font.bold = True
    r_link1.font.underline = True
    r_link1.font.color.rgb = ACCENT_CYAN
    r_link1.hyperlink.address = "https://harmony-household.vercel.app"

    p_c2 = tf11.add_paragraph()
    r2 = p_c2.add_run()
    r2.text = "2. Backend API Deployment: Render Web Service\n"
    r2.font.size = Pt(12)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT_VIOLET

    r_desc2 = p_c2.add_run()
    r_desc2.text = "   • Node.js/Express environment with automatic health checks (/api/health).\n   • Direct API Endpoint: "
    r_desc2.font.size = Pt(11)
    r_desc2.font.color.rgb = TEXT_WHITE

    r_link2 = p_c2.add_run()
    r_link2.text = "https://harmony-api.onrender.com/api/health\n"
    r_link2.font.size = Pt(12)
    r_link2.font.bold = True
    r_link2.font.underline = True
    r_link2.font.color.rgb = ACCENT_CYAN
    r_link2.hyperlink.address = "https://harmony-api.onrender.com/api/health"

    p_c3 = tf11.add_paragraph()
    r3 = p_c3.add_run()
    r3.text = "3. Version Control & Source Code Repository: GitHub\n"
    r3.font.size = Pt(12)
    r3.font.bold = True
    r3.font.color.rgb = ACCENT_AMBER

    r_desc3 = p_c3.add_run()
    r_desc3.text = "   • Continuous Integration from main branch with verified commit history.\n   • Repository URL: "
    r_desc3.font.size = Pt(11)
    r_desc3.font.color.rgb = TEXT_WHITE

    r_link3 = p_c3.add_run()
    r_link3.text = "https://github.com/lakshya212/HARMONY\n"
    r_link3.font.size = Pt(12)
    r_link3.font.bold = True
    r_link3.font.underline = True
    r_link3.font.color.rgb = ACCENT_CYAN
    r_link3.hyperlink.address = "https://github.com/lakshya212/HARMONY"

    add_footer(s11, 11)

    # ==========================================
    # SLIDE 12: FINAL PRESENTATION SUMMARY & DEFENSE
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_slide_background(s12)
    add_header(s12, "Final Presentation Summary & Viva Defense", category_text="MINOR PROJECT STATUS: 100% COMPLETE & DEPLOYED")

    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), title="FINAL DELIVERABLES READY")
    tb12_1 = s12.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.1))
    tf12_1 = tb12_1.text_frame
    tf12_1.word_wrap = True
    p = tf12_1.paragraphs[0]
    p.text = "Mandatory Submissions for Panel 2:\n\n[X] PDF of Project Along with Code:\n    Comprehensive documentation with Cannock Pvt Ltd certificate, DFDs, ERD, and Section 12 full source code appendix.\n\n[X] Project Deployment Link:\n    Production web app live on Vercel with resilient client fallback.\n\n[X] PPT for Presentation (Mandatory):\n    12-slide comprehensive presentation with live clickable links.\n\n[X] GitHub Source Code Repository:\n    Publicly accessible at github.com/lakshya212/HARMONY."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    add_card(s12, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9), title="PANEL 2 VIVA PRESENTATION", border_color=ACCENT_EMERALD)
    tb12_2 = s12.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.1))
    tf12_2 = tb12_2.text_frame
    tf12_2.word_wrap = True

    p = tf12_2.paragraphs[0]
    p.text = "Thank You!\n"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD

    p_sub = tf12_2.add_paragraph()
    p_sub.text = "HARMONY: Computational Transparency for Shared Living\n\nCandidate: Lakshya Shukla (Roll No: 240919106121)\nProject Guide: Mr. Nihal Kumar, Assistant Professor\nDepartment of Information Technology, IMS Noida\nAffiliated to Ch. Charan Singh University, Meerut\n\nOpen for Final Presentation Questions from Panel 2:\n• Mr. Nihal Kumar\n• Ms. Shaili Nigam"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = TEXT_WHITE

    add_footer(s12, 12)

    output_path = "BCA_508_Minor_Project_Presentation_HARMONY.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    create_deck()

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import win32com.client

from build_master_synopsis import (
    generate_master_document,
    set_cell_background,
    set_cell_borders,
    set_cell_margins,
    add_page_border
)

def add_code_block(doc, title, filename, code_content, description=None):
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(2)
    p_h.paragraph_format.keep_with_next = True
    
    r = p_h.add_run(f"{title}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0, 51, 102)
    
    r_fn = p_h.add_run(f"  [{filename}]")
    r_fn.font.size = Pt(9.5)
    r_fn.font.italic = True
    r_fn.font.color.rgb = RGBColor(100, 116, 139)

    if description:
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.keep_with_next = True
        r_d = p_desc.add_run(description)
        r_d.font.size = Pt(9)
        r_d.font.italic = True
        r_d.font.color.rgb = RGBColor(71, 85, 105)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F8FAFC")
    
    b_gray = {"val": "single", "sz": "4", "color": "CBD5E1"}
    b_blue = {"val": "single", "sz": "12", "color": "003366"}
    set_cell_borders(cell, top=b_gray, bottom=b_gray, left=b_blue, right=b_gray)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(9.5)
    
    lines = code_content.splitlines()
    for i, line in enumerate(lines, start=1):
        r_num = p.add_run(f"{i:3d}  ")
        r_num.font.name = "Consolas"
        r_num.font.size = Pt(7.5)
        r_num.font.color.rgb = RGBColor(148, 163, 184)
        
        r_text = p.add_run(line + ("\n" if i < len(lines) else ""))
        r_text.font.name = "Consolas"
        r_text.font.size = Pt(7.5)
        r_text.font.color.rgb = RGBColor(30, 41, 59)

def generate_report_with_code():
    print("Generating base academic documentation...")
    doc = generate_master_document()
    
    # Add page break for Section 12: Source Code Implementation
    doc.add_page_break()
    
    p_sec12 = doc.add_paragraph()
    p_sec12.paragraph_format.space_before = Pt(12)
    p_sec12.paragraph_format.space_after = Pt(6)
    r_12 = p_sec12.add_run("12. Source Code Implementation & System Listings")
    r_12.bold = True
    r_12.font.size = Pt(14)
    r_12.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph(
        "This section presents the complete documented source code implementation for the HARMONY Shared Living System. "
        "The project follows a modular, 3-tier decoupled architecture comprising the PostgreSQL relational data schema, "
        "the Express.js REST API service layer, and the React 18 client state management and user interface modules. "
        "All modules are strictly documented, type-safe, and hosted on GitHub under continuous integration."
    )

    # 1. Database Schema
    schema_path = "database/schema.sql"
    if os.path.exists(schema_path):
        with open(schema_path, "r", encoding="utf-8") as f:
            add_code_block(
                doc,
                "12.1 Database Relational DDL & Constraints",
                "database/schema.sql",
                f.read(),
                "PostgreSQL schema enforcing 3NF normal form, primary keys, foreign keys with CASCADE rules, and unique indices."
            )

    # 2. Backend Server & App
    server_path = "backend/src/server.js"
    if os.path.exists(server_path):
        with open(server_path, "r", encoding="utf-8") as f:
            add_code_block(
                doc,
                "12.2 Backend Server Entrypoint",
                "backend/src/server.js",
                f.read(),
                "Node.js HTTP listener initialization and environment configuration."
            )

    app_path = "backend/src/app.js"
    if os.path.exists(app_path):
        with open(app_path, "r", encoding="utf-8") as f:
            add_code_block(
                doc,
                "12.3 Express Application Configuration & Middleware",
                "backend/src/app.js",
                f.read(),
                "CORS security, Morgan logging, JSON parsing, API route mounting, and centralized error handling."
            )

    # 3. Backend Routes
    routes = [
        ("12.4 Household Management API Routes", "backend/src/routes/households.js", "Household creation, invite code resolution, and member association."),
        ("12.5 Chore Task Allocation API Routes", "backend/src/routes/tasks.js", "Task creation, effort point validation, and completion status toggling."),
        ("12.6 Shared Expense & Settlement API Routes", "backend/src/routes/expenses.js", "Shared bill recording, participant split calculations, and ledger balance queries.")
    ]
    for title, path, desc in routes:
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                add_code_block(doc, title, path, f.read(), desc)

    # 4. Frontend Context (The Heart of the System)
    ctx_path = "frontend/src/context/AppContext.jsx"
    if os.path.exists(ctx_path):
        with open(ctx_path, "r", encoding="utf-8") as f:
            add_code_block(
                doc,
                "12.7 Frontend State Management & Offline Resilient Store",
                "frontend/src/context/AppContext.jsx",
                f.read(),
                "React Context provider containing net balance computation, chore point leaderboard aggregation, and synchronized offline fallback."
            )

    # 5. Core Frontend Pages
    pages = [
        ("12.8 Roommate Interactive Dashboard Page", "frontend/src/pages/DashboardPage.jsx", "Main dashboard rendering real-time financial ledger, chore leaderboards, and pending task lists."),
        ("12.9 Shared Expense Management Module", "frontend/src/pages/ExpensesPage.jsx", "Expense creation modal, category filters, and participant split breakdown."),
        ("12.10 Chore Task Allocation Module", "frontend/src/pages/TasksPage.jsx", "Effort-weighted task cards, assignment selectors, and point badges."),
        ("12.11 Role-Based Admin Management Console", "frontend/src/pages/AdminPage.jsx", "Household administration, member management, invite code regeneration, and audit controls.")
    ]
    for title, path, desc in pages:
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                add_code_block(doc, title, path, f.read(), desc)

    # 6. Deployment & Build Configuration
    configs = [
        ("12.12 Vite Build Tooling Configuration", "frontend/vite.config.js", "Vite bundling options, React plugin integration, and local API proxy rules."),
        ("12.13 Vercel Edge SPA Rewrite Rules", "frontend/vercel.json", "Single-Page Application routing rewrite rules preventing 404 on direct URL navigation.")
    ]
    for title, path, desc in configs:
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                add_code_block(doc, title, path, f.read(), desc)

    out_docx = "BCA_508_Minor_Project_Final_Documentation_and_Code.docx"
    doc.save(out_docx)
    print(f"Saved complete document with code to: {out_docx}")
    
    # Convert to PDF
    out_pdf = "BCA_508_Minor_Project_Final_Documentation_and_Code.pdf"
    print(f"Converting {out_docx} to {out_pdf} via Word COM...")
    word = win32com.client.Dispatch("Word.Application")
    try:
        w_doc = word.Documents.Open(os.path.abspath(out_docx))
        w_doc.SaveAs(os.path.abspath(out_pdf), FileFormat=17) # 17 = wdFormatPDF
        w_doc.Close()
        print(f"Successfully created: {out_pdf}")
    finally:
        word.Quit()

def generate_standalone_code_document():
    print("Generating standalone Source Code Listing document...")
    doc = docx.Document()
    
    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        add_page_border(section)

    # Title Page
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(40)
    p_title.paragraph_format.space_after = Pt(12)
    
    r1 = p_title.add_run("HARMONY\n")
    r1.bold = True
    r1.font.size = Pt(26)
    r1.font.color.rgb = RGBColor(0, 51, 102)

    r2 = p_title.add_run("ROOMMATE CONFLICT RESOLUTION AND SHARED LIVING MANAGEMENT SYSTEM\n\n")
    r2.bold = True
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(16, 185, 129)

    r3 = p_title.add_run("COMPLETE SOURCE CODE LISTING & IMPLEMENTATION SPECIFICATIONS\n\n")
    r3.bold = True
    r3.font.size = Pt(14)
    r3.font.color.rgb = RGBColor(71, 85, 105)

    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course.paragraph_format.space_after = Pt(28)
    r_c = p_course.add_run(
        "Submitted in partial fulfillment of the requirements for the award of degree of\n"
        "BACHELOR OF COMPUTER APPLICATIONS (BCA)\n"
        "BCA 5th Semester | Course Code: BCA-508 (Minor Project)"
    )
    r_c.font.size = Pt(11)

    p_cand = doc.add_paragraph()
    p_cand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cand.paragraph_format.space_after = Pt(36)
    r_cand = p_cand.add_run(
        "Submitted By:\n"
        "Lakshya Shukla\n"
        "University Roll No: 240919106121 (BCA 5th Sem, Sec: B)\n"
        "IMS Noida, Affiliated to CCS University, Meerut\n\n"
        "Project Guide: Mr. Nihal Kumar, Assistant Professor\n"
        "Review Panel: Panel 2 (Room DG4)\n"
        "Date of Submission: 24th October 2026"
    )
    r_cand.font.size = Pt(11)

    doc.add_page_break()

    # Table of Files
    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_before = Pt(12)
    p_toc.paragraph_format.space_after = Pt(8)
    r_t = p_toc.add_run("INDEX OF IMPLEMENTATION MODULES")
    r_t.bold = True
    r_t.font.size = Pt(14)
    r_t.font.color.rgb = RGBColor(0, 51, 102)

    files_to_include = [
        ("Database Relational Schema (3NF DDL)", "database/schema.sql", "PostgreSQL tables, foreign key constraints, indexes"),
        ("Backend Server Listener", "backend/src/server.js", "Express HTTP server bootstrap and port listener"),
        ("Express Application & Middleware", "backend/src/app.js", "CORS, JSON body parser, route configuration, error handler"),
        ("Household API Controller", "backend/src/routes/households.js", "Household generation and code resolution endpoints"),
        ("Chore Engine API Controller", "backend/src/routes/tasks.js", "Task creation, effort point validation, toggle status"),
        ("Expense Splitting API Controller", "backend/src/routes/expenses.js", "Shared expense recording and participant split logic"),
        ("Frontend State Context & Failover Engine", "frontend/src/context/AppContext.jsx", "Client state management, net balance engine, offline fallback"),
        ("Roommate Live Dashboard Component", "frontend/src/pages/DashboardPage.jsx", "Interactive analytics, financial ledger, task boards"),
        ("Expense Management Component", "frontend/src/pages/ExpensesPage.jsx", "Shared bill entry modal, split previews, transaction history"),
        ("Chore Task Allocation Component", "frontend/src/pages/TasksPage.jsx", "Task list, effort-weighted assignment, status badges"),
        ("Household Admin Control Component", "frontend/src/pages/AdminPage.jsx", "Role-based controls, member eviction, ledger reset"),
        ("Vite Build Tooling Configuration", "frontend/vite.config.js", "Vite plugins, asset bundling, API reverse proxy"),
        ("Vercel Edge SPA Routing Rules", "frontend/vercel.json", "Production SPA URL rewrite configuration")
    ]

    table = doc.add_table(rows=len(files_to_include) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["S.No", "Module / File Name", "Architectural Role"]
    col_widths = [Inches(0.6), Inches(2.8), Inches(3.4)]

    for c_i, h in enumerate(headers):
        cell = table.cell(0, c_i)
        cell.width = col_widths[c_i]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_i, (m_title, m_path, m_desc) in enumerate(files_to_include, start=1):
        c0 = table.cell(r_i, 0)
        c1 = table.cell(r_i, 1)
        c2 = table.cell(r_i, 2)
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        c2.width = col_widths[2]
        
        bg = "F8FAFC" if r_i % 2 == 1 else "FFFFFF"
        for c in [c0, c1, c2]:
            set_cell_background(c, bg)
            set_cell_margins(c, 50, 50, 70, 70)
            
        c0.paragraphs[0].add_run(str(r_i)).font.size = Pt(9)
        r_fn = c1.paragraphs[0].add_run(f"{m_title}\n")
        r_fn.bold = True
        r_fn.font.size = Pt(9)
        c1.paragraphs[0].add_run(m_path).font.size = Pt(8)
        c2.paragraphs[0].add_run(m_desc).font.size = Pt(8.5)

    doc.add_page_break()

    # Append each code file
    for s_no, (m_title, m_path, m_desc) in enumerate(files_to_include, start=1):
        if os.path.exists(m_path):
            with open(m_path, "r", encoding="utf-8") as f:
                content = f.read()
                add_code_block(doc, f"{s_no}. {m_title}", m_path, content, m_desc)

    standalone_docx = "BCA_508_HARMONY_Source_Code_Listing.docx"
    doc.save(standalone_docx)
    print(f"Saved standalone code document to: {standalone_docx}")

    standalone_pdf = "BCA_508_HARMONY_Source_Code_Listing.pdf"
    print(f"Converting {standalone_docx} to {standalone_pdf} via Word COM...")
    word = win32com.client.Dispatch("Word.Application")
    try:
        w_doc = word.Documents.Open(os.path.abspath(standalone_docx))
        w_doc.SaveAs(os.path.abspath(standalone_pdf), FileFormat=17)
        w_doc.Close()
        print(f"Successfully created: {standalone_pdf}")
    finally:
        word.Quit()

if __name__ == "__main__":
    generate_report_with_code()
    generate_standalone_code_document()
